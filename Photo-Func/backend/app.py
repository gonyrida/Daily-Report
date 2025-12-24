import os
import io
import json
from io import BytesIO
from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

app = Flask(__name__)
CORS(app)

# =======================
# CONFIG
# =======================
TEMPLATE_FILE = os.path.join(os.path.dirname(__file__), "template.xlsx")

# Columns (1-based)
DATA_COLUMNS = [2, 4]    # B, D
FOOTER_COLUMNS = [2, 4]  # B, D

# Layout
TOP_PADDING = 1
BOTTOM_PADDING = 1
ENTRY_ROW_SPACING = TOP_PADDING + 1 + BOTTOM_PADDING + 1  # padding + image + padding + footer
START_DATA_ROW = 7        # first image row

# =======================
# HELPERS
# =======================
def write_to_merged_safe(ws, row, col, value):
    """
    Write to merged cell safely.
    If target is inside a merged range, write to top-left cell.
    """
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            ws.cell(row=merged.min_row, column=merged.min_col).value = value
            return
    ws.cell(row=row, column=col).value = value


def copy_row(ws, src_row, tgt_row):
    ws.insert_rows(tgt_row)

    for col in range(1, ws.max_column + 1):
        src_cell = ws.cell(row=src_row, column=col)
        tgt_cell = ws.cell(row=tgt_row, column=col)

        tgt_cell.value = src_cell.value
        if src_cell.has_style:
            tgt_cell._style = copy(src_cell._style)

    ws.row_dimensions[tgt_row].height = ws.row_dimensions[src_row].height


def copy_merged_cells(ws, src_start, src_end, offset):
    for merged in list(ws.merged_cells.ranges):
        if src_start <= merged.min_row <= src_end:
            ws.merge_cells(
                start_row=merged.min_row + offset,
                start_column=merged.min_col,
                end_row=merged.max_row + offset,
                end_column=merged.max_col,
            )

# =======================
# ROUTES
# =======================
@app.route("/export", methods=["POST"])
def export_excel():
    sections = json.loads(request.form["data"])
    files = request.files

    wb = load_workbook(TEMPLATE_FILE)
    ws = wb.active

    ENTRY_BLOCK_START = 6   # top padding row
    ENTRY_BLOCK_END = 9     # footer row
    ENTRY_HEIGHT = ENTRY_BLOCK_END - ENTRY_BLOCK_START + 1

    current_row = ENTRY_BLOCK_START
    image_cache = {}

    for section in sections:
        # ---- SECTION HEADER ----
        ws.cell(row=current_row - 1, column=2).value = section.get("title")

        first_entry = True

        for entry in section.get("entries", []):
            if not first_entry:
                offset = current_row - ENTRY_BLOCK_START

                for r in range(ENTRY_BLOCK_START, ENTRY_BLOCK_END + 1):
                    copy_row(ws, r, current_row)

                copy_merged_cells(
                    ws,
                    ENTRY_BLOCK_START,
                    ENTRY_BLOCK_END,
                    offset
                )

            image_row = current_row + 1
            footer_row = current_row + 3

            # ---- IMAGES ----
            for idx, img_key in enumerate(entry.get("images", {})):
                file_obj = files.get(img_key)
                if not file_obj:
                    continue

                if img_key not in image_cache:
                    image_cache[img_key] = file_obj.read()

                img = XLImage(BytesIO(image_cache[img_key]))
                col_letter = get_column_letter(DATA_COLUMNS[idx])
                ws.add_image(img, f"{col_letter}{image_row}")

            # ---- FOOTERS ----
            for idx, text in enumerate(entry.get("footers", ["", ""])):
                write_to_merged_safe(ws, footer_row, FOOTER_COLUMNS[idx], text)

            current_row += ENTRY_HEIGHT
            first_entry = False

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="output.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

if __name__ == "__main__":
    app.run(debug=True)
