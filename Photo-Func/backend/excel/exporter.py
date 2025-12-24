from io import BytesIO
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

from .config import (
    TEMPLATE_FILE,
    DATA_COLUMNS,
    FOOTER_COLUMNS,
    ENTRY_BLOCK_START,
    ENTRY_BLOCK_END,
    ENTRY_HEIGHT,
)
from .helpers import write_to_merged_safe
from .entry_block import copy_row, copy_merged_cells


def export_excel(sections, files):
    wb = load_workbook(TEMPLATE_FILE)
    ws = wb.active

    # We start at the Section Header row (Row 5 in your logic)
    current_row = ENTRY_BLOCK_START - 1 
    image_cache = {}

    # Define the template source rows
    TEMPLATE_HEADER_ROW = ENTRY_BLOCK_START - 1
    TEMPLATE_ENTRY_START = ENTRY_BLOCK_START
    TEMPLATE_ENTRY_END = ENTRY_BLOCK_END

    for section in sections:
        # 1. HANDLE SECTION HEADER (Row 5)
        # Copy the header row style/merge from template if not the very first row of the sheet
        if current_row != TEMPLATE_HEADER_ROW:
            copy_row(ws, TEMPLATE_HEADER_ROW, current_row)
            # Section headers are usually merged A-D
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        
        write_to_merged_safe(ws, current_row, DATA_COLUMNS[0], section.get("title"))
        current_row += 1 # Move to the first entry row of this section

        for entry_idx, entry in enumerate(section.get("entries", [])):
            # 2. HANDLE ENTRY BLOCK (Rows 6-9)
            # If it's not the original template position, we must clone the structure
            if current_row != TEMPLATE_ENTRY_START:
                offset = current_row - TEMPLATE_ENTRY_START
                
                # Copy rows 6 through 9 structure (Padding, Image Row, Padding, Footer)
                for r in range(TEMPLATE_ENTRY_START, TEMPLATE_ENTRY_END + 1):
                    copy_row(ws, r, current_row + (r - TEMPLATE_ENTRY_START))

                # Copy the specific merged cells for the footers (AB and CD)
                copy_merged_cells(ws, TEMPLATE_ENTRY_START, TEMPLATE_ENTRY_END, offset)

            # 3. INSERT IMAGES (Row 7 logic)
            # Based on your structure, image is 1 row below entry start (Row 6+1=7)
            image_row = current_row + 1
            for idx, img_key in enumerate(entry.get("images", {})):
                file_obj = files.get(img_key)
                if not file_obj or idx >= len(DATA_COLUMNS):
                    continue

                if img_key not in image_cache:
                    file_obj.stream.seek(0)
                    image_cache[img_key] = file_obj.read()

                img = XLImage(BytesIO(image_cache[img_key]))
                
                # OPTIONAL: Set a fixed size to prevent the "messy" overlap
                # img.width, img.height = 380, 240 

                col_letter = get_column_letter(DATA_COLUMNS[idx])
                ws.add_image(img, f"{col_letter}{image_row}")

            # 4. INSERT FOOTERS (Row 9 logic)
            # Footer is 3 rows below entry start (Row 6+3=9)
            footer_row = current_row + 3
            footers = entry.get("footers", ["", ""])
            for idx, text in enumerate(footers):
                if idx < len(FOOTER_COLUMNS):
                    write_to_merged_safe(ws, footer_row, FOOTER_COLUMNS[idx], text)

            # Move current_row to the start of the next entry
            current_row += ENTRY_HEIGHT

    return wb