# generators/excel/engine.py
from openpyxl import load_workbook

# Import from the new "common" location
from ..common.config import (
    TEMPLATE_FILE, DATA_COLUMNS, FOOTER_COLUMNS,
    ENTRY_BLOCK_START, ENTRY_BLOCK_END, ENTRY_HEIGHT,
)
from ..common.helpers import write_to_merged_safe

# Import from sibling files in the excel/ folder
from .templates import copy_row
from .images import process_and_insert_images
from .writer import prepare_entry_block, write_footers

def export_excel(sections, files):
    wb = load_workbook(TEMPLATE_FILE)
    ws = wb.active

    current_row = ENTRY_BLOCK_START - 1 
    image_cache = {}

    TEMPLATE_HEADER_ROW = ENTRY_BLOCK_START - 1
    TEMPLATE_ENTRY_START = ENTRY_BLOCK_START
    TEMPLATE_ENTRY_END = ENTRY_BLOCK_END

    for section in sections:
        # 1. HANDLE SECTION HEADER
        if current_row != TEMPLATE_HEADER_ROW:
            copy_row(ws, TEMPLATE_HEADER_ROW, current_row)
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        
        write_to_merged_safe(ws, current_row, DATA_COLUMNS[0], section.get("title"))
        current_row += 1 

        for entry in section.get("entries", []):
            # 2. CLONE STRUCTURE & SET ROW HEIGHT
            prepare_entry_block(ws, current_row, TEMPLATE_ENTRY_START, TEMPLATE_ENTRY_END)

            # 3. INSERT ANCHORED IMAGES
            process_and_insert_images(ws, entry, files, image_cache, current_row + 1, DATA_COLUMNS)

            # 4. INSERT FOOTERS
            write_footers(ws, current_row + 3, entry, FOOTER_COLUMNS)

            current_row += ENTRY_HEIGHT

    return wb