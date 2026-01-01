import io
import os
import openpyxl
from flask import Flask, request, send_file
from flask_cors import CORS

# Import our generator pieces
from generators.excel.writer import (
    fill_reference_sheet,
    fill_report_header,
    fill_activities,
    fill_team_tables,
    fill_material_machinery_tables
)

app = Flask(__name__)
CORS(app)

# --- ROUTE 1: REPORT ONLY (For Layout & Table Verification) ---
@app.route("/generate-report", methods=["POST"])
def generate_report():
    data = request.json 
    
    # Load the standard report template
    template_path = os.path.join(os.path.dirname(__file__), "templates", "template.xlsx")
    wb = openpyxl.load_workbook(template_path)
    ws = wb.worksheets[0] 

    # Execute Report Logic
    fill_report_header(ws, data)
    fill_activities(ws, data)
    
    # Tables & Shifting
    team_shift = fill_team_tables(ws, data)
    fill_material_machinery_tables(ws, data, shifted_offset=team_shift)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Report_Only_Verification.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

# --- ROUTE 2: REFERENCE ONLY (For Image Verification) ---
@app.route("/generate-reference", methods=["POST"])
def generate_reference():
    data = request.json 
    
    # Load the reference template (the one with the logo and reference header)
    template_path = os.path.join(os.path.dirname(__file__), "templates", "reference-template.xlsx")
    wb = openpyxl.load_workbook(template_path)
    ws = wb.worksheets[0] 

    # Only execute the photo logic
    reference_data = data.get("reference", [])
    if reference_data:
        fill_reference_sheet(ws, reference_data)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Reference_Only_Verification.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

# --- ROUTE 3: COMBINED (The Final Goal) ---
@app.route("/generate-combined", methods=["POST"])
def generate_combined():
    data = request.json 
    template_path = os.path.join(os.path.dirname(__file__), "templates", "template.xlsx")
    wb = openpyxl.load_workbook(template_path)
    ws = wb.worksheets[0] 

    fill_report_header(ws, data)
    fill_activities(ws, data)
    team_shift = fill_team_tables(ws, data)
    total_shift = fill_material_machinery_tables(ws, data, shifted_offset=team_shift)
    
    reference_data = data.get("reference", [])
    if reference_data:
        fill_reference_sheet(ws, reference_data, shifted_offset=total_shift)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Full_Combined_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5001, debug=True)