# generators/excel/engine.py
import os
from openpyxl import load_workbook
from .sheets.report import (
    fill_report_header, 
    fill_activities, 
    fill_team_tables, 
    fill_material_machinery_tables
)
from .sheets.reference import fill_reference_sheet

def generate_full_report(data, mode="report"):
    # 1. PATH SETUP
    # Goes up from generators/excel/ to project root
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    
    # 2. SEPARATE FILE LOGIC
    if mode == "report":
        template_name = "template.xlsx"
        template_path = os.path.join(base_dir, "templates", template_name)
        wb = load_workbook(template_path)
        ws = wb.worksheets[0]
        
        # Run ONLY Report Logic
        fill_report_header(ws, data)
        fill_activities(ws, data)
        team_shift = fill_team_tables(ws, data)
        fill_material_machinery_tables(ws, data, team_shift)
        
    elif mode == "reference":
        template_name = "reference-template.xlsx"
        template_path = os.path.join(base_dir, "templates", template_name)
        wb = load_workbook(template_path)
        ws = wb.worksheets[0]
        
        # 1. Run Reference Logic (This now includes our 4-entry counter + breaks)
        fill_reference_sheet(ws, data.get("reference", []))

        # 2. FINAL PRINT CALIBRATION
        # We must set fitToHeight to False (0) so Excel respects our MANUAL breaks
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0 
        
        # Optional: Force Portrait and A4 for consistency
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4

    return wb